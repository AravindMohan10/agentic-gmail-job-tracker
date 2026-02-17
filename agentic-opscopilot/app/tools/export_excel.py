# app/tools/export_excel.py
"""
Export application records to an Excel (.xlsx) file.
Reads from repository; db_path from env APPLICATIONS_DB_PATH.
"""
from __future__ import annotations

import os
from typing import Any, Dict, List, Optional

from app.repository.applications_repository import list_applications as _repo_list

_DEFAULT_DB_PATH = "applications.db"
_EXPORT_LIMIT = 100_000  # effectively all records for export


def _get_db_path() -> str:
    return os.environ.get("APPLICATIONS_DB_PATH", _DEFAULT_DB_PATH)


# Column order for the sheet
_COLUMNS = [
    "id",
    "message_id",
    "company",
    "role",
    "status",
    "stage",
    "event_date",
    "confidence",
    "raw_reason",
    "created_at",
    "updated_at",
]


def export_excel(
    *,
    path: str,
    sheet_name: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Write all application records to a .xlsx file.
    path must end with .xlsx (enforced by policy).
    """
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError(
            "openpyxl is required for export_excel. Install with: pip install openpyxl"
        ) from e

    db_path = _get_db_path()
    rows = _repo_list(
        db_path,
        limit=_EXPORT_LIMIT,
        status=None,
    )

    wb = openpyxl.Workbook()
    sheet = wb.active
    if sheet_name and sheet_name.strip():
        sheet.title = sheet_name.strip()[:31]  # Excel sheet name max 31 chars

    for col_idx, header in enumerate(_COLUMNS, start=1):
        sheet.cell(row=1, column=col_idx, value=header)

    for row_idx, record in enumerate(rows, start=2):
        for col_idx, key in enumerate(_COLUMNS, start=1):
            value = record.get(key)
            sheet.cell(row=row_idx, column=col_idx, value=value)

    wb.save(path)
    return {"path": path, "rows_exported": len(rows)}
